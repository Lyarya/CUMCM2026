"""End-to-end runner for Problem 1 deterministic dispatch."""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.common.paths import PROCESSED_DATA_DIR, PROJECT_ROOT, problem_results_dir
from src.problem1.evaluate import build_summary, compute_baseline, validate_dispatch
from src.problem1.model import DispatchParameters, solve_deterministic_dispatch
from src.problem1.visualize import create_q1_figures


INPUT_PATH = PROCESSED_DATA_DIR / "C题" / "problem1_day.csv"
TEMPLATE_PATH = PROJECT_ROOT / "data" / "raw" / "C题" / "附件" / "附件5" / "result1.xlsx"
OUTPUT_DIR = problem_results_dir(1)
TABLE_DIR = OUTPUT_DIR / "tables"
RESULT1_PATH = OUTPUT_DIR / "result1.xlsx"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def export_result1(
    dispatch: pd.DataFrame,
    output_path: Path = RESULT1_PATH,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    """Fill the official two-sheet result1 template without changing its layout."""
    if len(dispatch) != 144:
        raise ValueError("The official result1 template requires 144 dispatch periods")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = openpyxl.load_workbook(output_path)
    if workbook.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError(f"Unexpected result1 sheets: {workbook.sheetnames}")

    purchase_sheet = workbook["计划购电量"]
    if purchase_sheet.max_row != 145:
        raise ValueError("The official purchase sheet must contain 144 data rows")
    for offset, value in enumerate(dispatch["grid_purchase_kwh"], start=2):
        cell = purchase_sheet.cell(row=offset, column=2)
        cell.value = float(value)
        cell.number_format = "0.000000"

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        rows = dispatch.iloc[block * 24 : (block + 1) * 24]
        storage_sheet.cell(row=block + 2, column=2, value=float(rows["charge_kwh"].sum())).number_format = "0.000000"
        storage_sheet.cell(row=block + 2, column=3, value=float(rows["discharge_kwh"].sum())).number_format = "0.000000"
    storage_sheet["E2"] = float(dispatch["storage_start_kwh"].iloc[0])
    storage_sheet["E3"] = float(dispatch["storage_end_kwh"].iloc[-1])
    # Keep the official column width while displaying the endpoint values fully.
    storage_sheet["E2"].number_format = "0.000"
    storage_sheet["E3"].number_format = "0.000"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)

    check = openpyxl.load_workbook(output_path, data_only=True)
    if check.sheetnames != ["计划购电量", "充放电量"]:
        raise AssertionError("result1 sheet structure changed during export")
    values = [check["计划购电量"].cell(row=row, column=2).value for row in range(2, 146)]
    if len(values) != 144 or any(value is None for value in values):
        raise AssertionError("result1 purchase sheet is incomplete")
    if abs(float(check["充放电量"]["E2"].value) - 6_000.0) > 1e-6:
        raise AssertionError("result1 initial storage energy is incorrect")
    if abs(float(check["充放电量"]["E3"].value) - 6_000.0) > 1e-6:
        raise AssertionError("result1 terminal storage energy is incorrect")
    return output_path


def main() -> None:
    raw_paths = sorted((PROJECT_ROOT / "data" / "raw" / "C题" / "附件").glob("附件[1-4].xlsx"))
    raw_hashes_before = {path.name: _sha256(path) for path in raw_paths}
    processed_hash_before = _sha256(INPUT_PATH)
    data = pd.read_csv(INPUT_PATH)
    parameters = DispatchParameters()
    result = solve_deterministic_dispatch(data, parameters)
    validation = validate_dispatch(result, parameters)
    baseline = compute_baseline(data)
    summary = build_summary(result, validation, baseline, parameters)

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    result.dispatch.to_csv(TABLE_DIR / "table_p1_dispatch.csv", index=False, float_format="%.10g")
    pd.DataFrame(
        [{"metric": key, "value": value} for key, value in summary.items()]
    ).to_csv(TABLE_DIR / "table_p1_summary.csv", index=False)
    (TABLE_DIR / "table_p1_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    export_result1(result.dispatch)
    figure_paths = create_q1_figures(result.dispatch)

    raw_hashes_after = {path.name: _sha256(path) for path in raw_paths}
    if raw_hashes_before != raw_hashes_after:
        raise AssertionError("Raw official workbooks changed during Q1 execution")
    if processed_hash_before != _sha256(INPUT_PATH):
        raise AssertionError("Stage 1A problem1_day.csv changed during Q1 execution")

    print("Q1 deterministic dispatch complete.")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"result1: {RESULT1_PATH} (144 x 2 purchase sheet; 6 x 5 storage summary)")
    print("figures:")
    for figure in figure_paths.values():
        print("  " + ", ".join(str(path) for path in figure.values()))


if __name__ == "__main__":
    main()
