"""Chronological runner and official result2 exporter for Problem 2."""

from __future__ import annotations

import argparse
from copy import copy
import hashlib
import json
from pathlib import Path
import shutil
import sys

import numpy as np
import openpyxl
import pandas as pd

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.common.paths import PROJECT_ROOT, problem_results_dir
from src.problem2.evaluate import build_daily_summary, validate_q2_day
from src.problem2.forecast_interface import (
    FINAL_AUDIT_PATH,
    FORECAST_DIR,
    FORMAL_FORECAST_PATH,
    SCENARIO_PATH,
    audit_q2_handoff_integrity,
    get_q2_day_inputs,
)
from src.problem2.model import Q2DayResult, Q2DispatchParameters, solve_expected_cost_dispatch


FORMAL_START = "2025-02-01"
FORMAL_END = "2025-12-31"
OFFICIAL_INITIAL_ENERGY_KWH = 6_000.0
TABLE3_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
OUTPUT_DIR = problem_results_dir(2)
TABLE_DIR = OUTPUT_DIR / "tables"
TEMPLATE_PATH = PROJECT_ROOT / "data" / "raw" / "C题" / "附件" / "附件5" / "result2.xlsx"
RESULT2_PATH = OUTPUT_DIR / "result2.xlsx"
PROTECTED_FORECAST_PATHS = (
    FORMAL_FORECAST_PATH,
    SCENARIO_PATH,
    FORECAST_DIR / "q2_scenario_manifest.csv",
    FORECAST_DIR / "q2_forecast_metrics.csv",
    FORECAST_DIR / "q2_model_decision.json",
    FORECAST_DIR / "q2_leakage_audit.json",
    FINAL_AUDIT_PATH,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _capture_row_style(sheet, source_row: int, column_count: int) -> tuple[object, list[tuple[object, str, object, object]]]:
    cells = []
    for column in range(1, column_count + 1):
        source = sheet.cell(source_row, column)
        cells.append(
            (
                copy(source._style),
                source.number_format,
                copy(source.alignment),
                copy(source.protection),
            )
        )
    return sheet.row_dimensions[source_row].height, cells


def _apply_row_style(sheet, style, target_row: int) -> None:
    height, cells = style
    sheet.row_dimensions[target_row].height = height
    for column, (cell_style, number_format, alignment, protection) in enumerate(cells, start=1):
        target = sheet.cell(target_row, column)
        target._style = copy(cell_style)
        target.number_format = number_format
        target.alignment = copy(alignment)
        target.protection = copy(protection)


def _interval_label(slot: int) -> str:
    start_minutes = int(slot) * 10
    end_minutes = (int(slot) + 1) * 10
    start_hour, start_minute = divmod(start_minutes % (24 * 60), 60)
    end_hour, end_minute = divmod(end_minutes % (24 * 60), 60)
    suffix = "+1" if end_minutes >= 24 * 60 else ""
    return f"{start_hour}:{start_minute:02d}-{end_hour}:{end_minute:02d}{suffix}"


def settle_realized_day(
    plan: pd.DataFrame,
    actual: pd.DataFrame,
    initial_energy_kwh: float,
    *,
    minimum_energy_kwh: float = 1_200.0,
    maximum_energy_kwh: float = 10_800.0,
    charge_efficiency: float = 0.9,
    discharge_efficiency: float = 0.9,
    dt_hours: float = 1.0 / 6.0,
) -> tuple[pd.DataFrame, float]:
    """Replay one day with planned battery actions as physical upper bounds."""

    frame = plan.copy().sort_values("slot", kind="stable").reset_index(drop=True)
    actual = actual.copy().sort_values("datetime", kind="stable").reset_index(drop=True)
    if len(frame) != 144 or len(actual) != 144:
        raise AssertionError("a Q2 day must contain exactly 144 intervals")

    energy = float(initial_energy_kwh)
    realized_rows: list[dict[str, float]] = []
    for index, row in frame.iterrows():
        load_kw = float(actual.at[index, "actual_load"])
        pv_kw = max(float(actual.at[index, "actual_generation"]), 0.0)
        planned_grid_kw = max(float(row["planned_grid_kw"]), 0.0)
        planned_charge_kw = max(float(row["planned_charge_limit_kw"]), 0.0)
        planned_discharge_kw = max(float(row["planned_discharge_limit_kw"]), 0.0)

        charge_headroom_kw = max(
            (maximum_energy_kwh - energy) / (charge_efficiency * dt_hours), 0.0
        )
        actual_charge_kw = min(planned_charge_kw, charge_headroom_kw)

        demand_after_pv_kw = load_kw + actual_charge_kw - pv_kw
        available_discharge_kw = max(
            (energy - minimum_energy_kwh) * discharge_efficiency / dt_hours, 0.0
        )
        actual_discharge_kw = min(
            planned_discharge_kw,
            max(demand_after_pv_kw, 0.0),
            available_discharge_kw,
        )
        net_after_battery_kw = demand_after_pv_kw - actual_discharge_kw
        grid_used_kw = min(planned_grid_kw, max(net_after_battery_kw, 0.0))
        emergency_kw = max(net_after_battery_kw - grid_used_kw, 0.0)
        pv_spill_kw = max(-net_after_battery_kw, 0.0)

        next_energy = (
            energy
            + charge_efficiency * actual_charge_kw * dt_hours
            - actual_discharge_kw * dt_hours / discharge_efficiency
        )
        if next_energy < minimum_energy_kwh - 2e-6 or next_energy > maximum_energy_kwh + 2e-6:
            raise AssertionError("realized SOC left the permitted range")
        next_energy = min(max(next_energy, minimum_energy_kwh), maximum_energy_kwh)
        realized_rows.append(
            {
                "actual_load_kw": load_kw,
                "actual_pv_kw": pv_kw,
                "actual_charge_kw": actual_charge_kw,
                "actual_discharge_kw": actual_discharge_kw,
                "actual_storage_start_kwh": energy,
                "actual_storage_end_kwh": next_energy,
                "realized_planned_grid_used_kw": grid_used_kw,
                "realized_unused_planned_grid_kw": planned_grid_kw - grid_used_kw,
                "realized_emergency_kw": emergency_kw,
                "realized_pv_spill_kw": pv_spill_kw,
            }
        )
        energy = next_energy

    frame = pd.concat([frame, pd.DataFrame(realized_rows)], axis=1)
    for column in (
        "actual_charge",
        "actual_discharge",
        "realized_planned_grid_used",
        "realized_unused_planned_grid",
        "realized_emergency",
        "realized_pv_spill",
    ):
        frame[f"{column}_kwh"] = frame[f"{column}_kw"] * dt_hours
    frame["realized_emergency_cost_yuan"] = (
        5.0 * frame["price_yuan_per_kwh"] * frame["realized_emergency_kwh"]
    )
    return frame, energy


def attach_realized_outcomes(
    daily: pd.DataFrame, intervals: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Replay existing plans sequentially with actual SOC carried across days."""

    actual = pd.read_csv(
        FORMAL_FORECAST_PATH,
        usecols=["operating_date", "datetime", "actual_load", "actual_generation"],
        parse_dates=["datetime"],
    ).sort_values(["operating_date", "datetime"], kind="stable")
    frame = intervals.copy().sort_values(["date", "timestamp"], kind="stable").reset_index(drop=True)
    actual = actual.loc[actual["operating_date"].astype(str).isin(frame["date"].astype(str).unique())]
    actual = actual.reset_index(drop=True)
    if len(frame) != len(actual):
        raise AssertionError("dispatch and Appendix 2 actual observations have different lengths")
    if not np.array_equal(frame["date"].astype(str), actual["operating_date"].astype(str)):
        raise AssertionError("dispatch dates do not align with Appendix 2 actual observations")
    if not np.array_equal(
        pd.to_datetime(frame["timestamp"]).to_numpy(dtype="datetime64[ns]"),
        actual["datetime"].to_numpy(dtype="datetime64[ns]"),
    ):
        raise AssertionError("dispatch timestamps do not align with Appendix 2 actual observations")

    settled_days: list[pd.DataFrame] = []
    realized_end_by_date: dict[str, float] = {}
    carried_energy = float(daily.iloc[0]["initial_energy_kwh"])
    for date, planned_day in frame.groupby("date", sort=True):
        actual_day = actual.loc[actual["operating_date"].astype(str) == str(date)]
        settled_day, carried_energy = settle_realized_day(
            planned_day, actual_day, carried_energy
        )
        settled_days.append(settled_day)
        realized_end_by_date[str(date)] = carried_energy
    frame = pd.concat(settled_days, ignore_index=True)

    realized_daily = frame.groupby("date", sort=True).agg(
        realized_emergency_energy_kwh=("realized_emergency_kwh", "sum"),
        realized_emergency_cost_yuan=("realized_emergency_cost_yuan", "sum"),
        realized_planned_grid_used_energy_kwh=("realized_planned_grid_used_kwh", "sum"),
        realized_unused_planned_grid_energy_kwh=("realized_unused_planned_grid_kwh", "sum"),
        realized_pv_spill_energy_kwh=("realized_pv_spill_kwh", "sum"),
        actual_charge_energy_kwh=("actual_charge_kwh", "sum"),
        actual_discharge_energy_kwh=("actual_discharge_kwh", "sum"),
    )
    daily_frame = daily.copy()
    for column in realized_daily.columns:
        daily_frame[column] = daily_frame["date"].map(realized_daily[column])
    if daily_frame[list(realized_daily.columns)].isna().any().any():
        raise AssertionError("realized Q2 daily outcomes contain missing values")
    daily_frame["realized_total_cost_yuan"] = (
        daily_frame["planned_purchase_cost_yuan"]
        + daily_frame["realized_emergency_cost_yuan"]
    )
    daily_frame["planned_final_energy_kwh"] = daily_frame["final_energy_kwh"]
    daily_frame["actual_final_energy_kwh"] = daily_frame["date"].map(realized_end_by_date)
    daily_frame["final_energy_kwh"] = daily_frame["actual_final_energy_kwh"]
    return daily_frame, frame


def write_table3_outputs(intervals: pd.DataFrame) -> tuple[Path, Path]:
    """Write the four-date emergency-purchase table in the official Table 3 form."""

    selected: dict[str, pd.DataFrame] = {}
    for target in TABLE3_DATES:
        rows = intervals.loc[
            intervals["date"] == target,
            ["slot", "realized_emergency_kwh"],
        ].sort_values("slot")
        if len(rows) != 144:
            raise AssertionError(f"{target} does not contain all 144 Table 3 intervals")
        selected[target] = rows.reset_index(drop=True)
    row_count = 144
    table = pd.DataFrame(index=range(row_count))
    display_dates = [
        f"{pd.Timestamp(target).year}.{pd.Timestamp(target).month}.{pd.Timestamp(target).day}"
        for target in TABLE3_DATES
    ]
    for target, display_date in zip(TABLE3_DATES, display_dates):
        rows = selected[target]
        table[f"{display_date}_时间段"] = [
            _interval_label(int(rows.at[index, "slot"]))
            for index in range(row_count)
        ]
        table[f"{display_date}_购电量_kWh"] = [
            float(rows.at[index, "realized_emergency_kwh"])
            for index in range(row_count)
        ]

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = TABLE_DIR / "table_p2_table3_emergency.csv"
    tex_path = TABLE_DIR / "table_p2_table3_emergency.tex"
    table.to_csv(csv_path, index=False, float_format="%.6f")
    lines = [
        "\\begin{table}[H]",
        "\\centering",
        "\\scriptsize",
        "\\setlength{\\tabcolsep}{2.5pt}",
        "\\caption{微网在指定日期的紧急购电量}",
        "\\label{tab:p2-table3-emergency}",
        "\\begin{tabular}{@{}crcrcrcr@{}}",
        "\\toprule",
        " & ".join(f"\\multicolumn{{2}}{{c}}{{{date}}}" for date in display_dates) + " \\\\",
        "\\cmidrule(lr){1-2}\\cmidrule(lr){3-4}\\cmidrule(lr){5-6}\\cmidrule(lr){7-8}",
        " & ".join(["时间段 & 购电量"] * 4) + " \\\\",
        "\\midrule",
    ]
    for row_index in range(6):
        cells: list[str] = []
        for target in TABLE3_DATES:
            rows = selected[target]
            cells.extend(
                [
                    _interval_label(int(rows.at[row_index, "slot"])).replace("-", "--"),
                    f"{float(rows.at[row_index, 'realized_emergency_kwh']):.3f}",
                ]
            )
        lines.append(" & ".join(cells) + " \\\\")
    lines.extend(
        [
            "\\multicolumn{2}{c}{$\\vdots$} & \\multicolumn{2}{c}{$\\vdots$} & "
            "\\multicolumn{2}{c}{$\\vdots$} & \\multicolumn{2}{c}{$\\vdots$} \\\\",
            "\\bottomrule",
            "\\end{tabular}",
            "\\par\\vspace{2pt}",
            "\\parbox{0.88\\textwidth}{\\footnotesize 注：购电量单位为 kWh。因篇幅限制，此处仅展示各指定日期前 6 个时段；全部 144 个 10 分钟时段（含紧急购电量为 0 的时段）已完整保存在附件 \\texttt{result2.xlsx} 中。}",
            "\\end{table}",
        ]
    )
    lines.append("")
    tex_path.write_text("\n".join(lines), encoding="utf-8")
    return csv_path, tex_path


def export_result2(
    daily: pd.DataFrame,
    intervals: pd.DataFrame,
    output_path: Path = RESULT2_PATH,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    """Fill the three-sheet official result2 workbook without editing its source."""

    dates = pd.to_datetime(daily["date"]).dt.normalize()
    if len(daily) != 334 or dates.iloc[0] != pd.Timestamp(FORMAL_START):
        raise ValueError("official result2 export requires all 334 formal days")
    if dates.iloc[-1] != pd.Timestamp(FORMAL_END):
        raise ValueError("official result2 export must end on 2025-12-31")
    if len(intervals) != 334 * 144:
        raise ValueError("official result2 export requires 144 intervals per formal day")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = openpyxl.load_workbook(output_path)
    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"unexpected result2 sheets: {workbook.sheetnames}")

    purchase_sheet = workbook["计划购电量"]
    if purchase_sheet.max_row != 335 or purchase_sheet.max_column != 147:
        raise ValueError("official planned-purchase sheet shape changed")
    for slot in range(1, 145):
        purchase_sheet.cell(1, slot + 1, _interval_label(slot))
    for day_index, (_, row) in enumerate(daily.iterrows(), start=2):
        day_frame = intervals.loc[intervals["date"] == row["date"]].sort_values("slot")
        if len(day_frame) != 144:
            raise AssertionError(f"{row['date']} does not contain 144 dispatch intervals")
        for slot, value in enumerate(day_frame["planned_grid_kwh"], start=2):
            cell = purchase_sheet.cell(day_index, slot, float(value))
            cell.number_format = "0.000000"
        purchase_sheet.cell(day_index, 146, float(row["planned_purchase_energy_kwh"])).number_format = "0.000000"
        purchase_sheet.cell(day_index, 147, float(row["planned_purchase_cost_yuan"])).number_format = "0.000000"

    storage_sheet = workbook["充放电量"]
    storage_styles = [_capture_row_style(storage_sheet, 2 + block, 6) for block in range(6)]
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    time_blocks = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    target_row = 2
    for _, day_row in daily.iterrows():
        day_frame = intervals.loc[intervals["date"] == day_row["date"]].sort_values("slot")
        for block, time_label in enumerate(time_blocks):
            _apply_row_style(storage_sheet, storage_styles[block], target_row)
            block_frame = day_frame.iloc[block * 24 : (block + 1) * 24]
            storage_sheet.cell(target_row, 1, pd.Timestamp(day_row["date"]).to_pydatetime() if block == 0 else None)
            storage_sheet.cell(target_row, 2, time_label)
            storage_sheet.cell(target_row, 3, float(block_frame["actual_charge_kwh"].sum())).number_format = "0.000000"
            storage_sheet.cell(target_row, 4, float(block_frame["actual_discharge_kwh"].sum())).number_format = "0.000000"
            if block == 0:
                storage_sheet.cell(target_row, 5, "00:00")
                storage_sheet.cell(target_row, 6, float(day_row["actual_initial_energy_kwh"])).number_format = "0.000000"
            elif block == 1:
                storage_sheet.cell(target_row, 5, "24:00")
                storage_sheet.cell(target_row, 6, float(day_row["actual_final_energy_kwh"])).number_format = "0.000000"
            else:
                storage_sheet.cell(target_row, 5, None)
                storage_sheet.cell(target_row, 6, None)
            target_row += 1

    emergency_sheet = workbook["紧急购电量"]
    emergency_styles = [_capture_row_style(emergency_sheet, 2 + row, 3) for row in range(3)]
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    target_row = 2
    for _, day_row in daily.iterrows():
        day_frame = intervals.loc[intervals["date"] == day_row["date"]].sort_values("slot")
        if len(day_frame) != 144:
            raise AssertionError(f"{day_row['date']} does not contain 144 emergency intervals")
        for item_index, (_, interval_row) in enumerate(day_frame.iterrows()):
            _apply_row_style(emergency_sheet, emergency_styles[min(item_index, 2)], target_row)
            emergency_sheet.cell(
                target_row,
                1,
                pd.Timestamp(day_row["date"]).to_pydatetime() if item_index == 0 else None,
            )
            emergency_sheet.cell(target_row, 2, _interval_label(int(interval_row["slot"])))
            emergency_sheet.cell(target_row, 3, float(interval_row["realized_emergency_kwh"])).number_format = "0.000000"
            target_row += 1

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)

    check = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    if check.sheetnames != expected_sheets:
        raise AssertionError("result2 sheet structure changed during export")
    if check["计划购电量"].max_row != 335 or check["计划购电量"].max_column != 147:
        raise AssertionError("result2 planned-purchase sheet shape is incorrect")
    return output_path


def run_chronological(
    dates: pd.DatetimeIndex,
    *,
    initial_energy_kwh: float = OFFICIAL_INITIAL_ENERGY_KWH,
    parameters: Q2DispatchParameters | None = None,
    progress: bool = False,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
]:
    """Solve and settle dates in order, carrying realized final SOC forward."""

    parameters = parameters or Q2DispatchParameters()
    carried_energy = float(initial_energy_kwh)
    summaries: list[dict[str, object]] = []
    interval_frames: list[pd.DataFrame] = []
    emergency_days: list[np.ndarray] = []
    spill_days: list[np.ndarray] = []
    planned_grid_used_days: list[np.ndarray] = []
    unused_planned_grid_days: list[np.ndarray] = []
    source_days: list[np.ndarray] = []
    actual_observations = pd.read_csv(
        FORMAL_FORECAST_PATH,
        usecols=["operating_date", "datetime", "actual_load", "actual_generation"],
        parse_dates=["datetime"],
    )
    for day_index, target in enumerate(dates, start=1):
        inputs = get_q2_day_inputs(target, initial_energy=carried_energy)
        result: Q2DayResult = solve_expected_cost_dispatch(inputs, parameters)
        validation = validate_q2_day(inputs, result, parameters)
        summary = build_daily_summary(result, validation, inputs)
        actual_day = actual_observations.loc[
            actual_observations["operating_date"].astype(str) == str(target.date())
        ]
        settled_day, realized_end_energy = settle_realized_day(
            result.dispatch, actual_day, carried_energy
        )
        summary["planned_final_energy_kwh"] = summary["final_energy_kwh"]
        summary["actual_initial_energy_kwh"] = carried_energy
        summary["actual_final_energy_kwh"] = realized_end_energy
        summary["final_energy_kwh"] = realized_end_energy
        summary["realized_emergency_energy_kwh"] = float(
            settled_day["realized_emergency_kwh"].sum()
        )
        summary["realized_emergency_cost_yuan"] = float(
            settled_day["realized_emergency_cost_yuan"].sum()
        )
        summary["realized_total_cost_yuan"] = float(
            summary["planned_purchase_cost_yuan"]
            + summary["realized_emergency_cost_yuan"]
        )
        summary["realized_planned_grid_used_energy_kwh"] = float(
            settled_day["realized_planned_grid_used_kwh"].sum()
        )
        summary["realized_unused_planned_grid_energy_kwh"] = float(
            settled_day["realized_unused_planned_grid_kwh"].sum()
        )
        summary["realized_pv_spill_energy_kwh"] = float(
            settled_day["realized_pv_spill_kwh"].sum()
        )
        summary["actual_charge_energy_kwh"] = float(
            settled_day["actual_charge_kwh"].sum()
        )
        summary["actual_discharge_energy_kwh"] = float(
            settled_day["actual_discharge_kwh"].sum()
        )
        charge_clip = (
            settled_day["planned_charge_limit_kw"] - settled_day["actual_charge_kw"]
        ).clip(lower=0.0)
        discharge_clip = (
            settled_day["planned_discharge_limit_kw"]
            - settled_day["actual_discharge_kw"]
        ).clip(lower=0.0)
        summary["battery_clipping_energy_kwh"] = float(
            (charge_clip + discharge_clip).sum() * inputs.dt_hours
        )
        summary["battery_clipping_intervals"] = int(
            ((charge_clip + discharge_clip) > 1e-9).sum()
        )
        if summaries and abs(float(summary["initial_energy_kwh"]) - float(summaries[-1]["final_energy_kwh"])) > 2e-3:
            raise AssertionError(f"cross-day SOC discontinuity before {target.date()}")
        summaries.append(summary)
        interval_frames.append(settled_day)
        emergency_days.append(result.emergency_kw)
        spill_days.append(result.spill_kw)
        planned_grid_used_days.append(result.planned_grid_used_kw)
        unused_planned_grid_days.append(result.unused_planned_grid_kw)
        source_days.append(result.scenario_source_dates)
        carried_energy = realized_end_energy
        if progress and (day_index == 1 or day_index % 10 == 0 or day_index == len(dates)):
            print(
                f"Q2 progress: {day_index}/{len(dates)} days; "
                f"{target.date()} {result.status}; runtime={result.runtime_seconds:.3f}s",
                flush=True,
            )
    return (
        pd.DataFrame(summaries),
        pd.concat(interval_frames, ignore_index=True),
        np.stack(emergency_days),
        np.stack(spill_days),
        np.stack(planned_grid_used_days),
        np.stack(unused_planned_grid_days),
        np.stack(source_days),
    )


def _aggregate_summary(daily: pd.DataFrame, intervals: pd.DataFrame) -> dict[str, object]:
    return {
        "q2_formulation": "expected-cost stochastic MILP",
        "solver": str(daily["solver"].iloc[0]),
        "scenario_count": 50,
        "formal_days": int(len(daily)),
        "optimal_days": int((daily["status"] == "Optimal").sum()),
        "infeasible_days": int((daily["status"] != "Optimal").sum()),
        "planned_purchase_cost_yuan": float(daily["planned_purchase_cost_yuan"].sum()),
        "expected_emergency_cost_yuan": float(daily["expected_emergency_cost_yuan"].sum()),
        "expected_total_cost_yuan": float(daily["expected_total_cost_yuan"].sum()),
        "realized_emergency_cost_yuan": float(daily["realized_emergency_cost_yuan"].sum()),
        "realized_total_cost_yuan": float(daily["realized_total_cost_yuan"].sum()),
        "planned_purchase_energy_kwh": float(daily["planned_purchase_energy_kwh"].sum()),
        "expected_emergency_energy_kwh": float(daily["expected_emergency_energy_kwh"].sum()),
        "expected_planned_grid_used_energy_kwh": float(
            daily["expected_planned_grid_used_energy_kwh"].sum()
        ),
        "expected_unused_planned_grid_energy_kwh": float(
            daily["expected_unused_planned_grid_energy_kwh"].sum()
        ),
        "realized_emergency_energy_kwh": float(daily["realized_emergency_energy_kwh"].sum()),
        "realized_planned_grid_used_energy_kwh": float(
            daily["realized_planned_grid_used_energy_kwh"].sum()
        ),
        "realized_unused_planned_grid_energy_kwh": float(
            daily["realized_unused_planned_grid_energy_kwh"].sum()
        ),
        "realized_pv_spill_energy_kwh": float(
            daily["realized_pv_spill_energy_kwh"].sum()
        ),
        "unabsorbed_discharge_energy_kwh": 0.0,
        "planned_charge_energy_kwh": float(daily["charge_energy_kwh"].sum()),
        "planned_discharge_energy_kwh": float(daily["discharge_energy_kwh"].sum()),
        "actual_charge_energy_kwh": float(daily["actual_charge_energy_kwh"].sum()),
        "actual_discharge_energy_kwh": float(daily["actual_discharge_energy_kwh"].sum()),
        "battery_clipping_energy_kwh": float(daily["battery_clipping_energy_kwh"].sum()),
        "battery_clipping_intervals": int(daily["battery_clipping_intervals"].sum()),
        "expected_spill_energy_kwh": float(daily["expected_spill_energy_kwh"].sum()),
        "minimum_energy_kwh": float(intervals[["actual_storage_start_kwh", "actual_storage_end_kwh"]].min().min()),
        "maximum_energy_kwh": float(intervals[["actual_storage_start_kwh", "actual_storage_end_kwh"]].max().max()),
        "final_battery_energy_kwh": float(daily["final_energy_kwh"].iloc[-1]),
        "maximum_power_balance_residual_kw": float(daily["maximum_power_balance_residual_kw"].max()),
        "maximum_soc_residual_kwh": float(daily["maximum_soc_residual_kwh"].max()),
        "total_solver_runtime_seconds": float(daily["runtime_seconds"].sum()),
        "daily_reset_to_6000": False,
        "terminal_treatment": (
            "decreasing piecewise-linear continuation value; "
            "no daily terminal target or reset"
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--representative", action="store_true", help="solve four representative days only")
    args = parser.parse_args()
    dates = (
        pd.to_datetime(["2025-02-01", "2025-04-15", "2025-07-15", "2025-10-15"])
        if args.representative
        else pd.date_range(FORMAL_START, FORMAL_END, freq="D")
    )
    protected_before = {str(path): _sha256(path) for path in PROTECTED_FORECAST_PATHS}
    (
        daily,
        intervals,
        emergency,
        spill,
        planned_grid_used,
        unused_planned_grid,
        source_dates,
    ) = run_chronological(
        pd.DatetimeIndex(dates), progress=True
    )
    # Realized settlement and SOC carry are performed inside run_chronological.
    protected_after = {str(path): _sha256(path) for path in PROTECTED_FORECAST_PATHS}
    if protected_before != protected_after:
        raise AssertionError("a frozen Stage 2A forecast/scenario artifact changed")
    audit = audit_q2_handoff_integrity()
    if not all(
        bool(audit[key])
        for key in (
            "protected_artifact_hashes_unchanged",
            "q2_scenario_hash_unchanged",
            "raw_data_hashes_unchanged",
            "paired_scenarios",
            "leakage_audit",
        )
    ):
        raise AssertionError(f"Stage 2B handoff audit failed: {audit}")

    if args.representative:
        print(json.dumps(_aggregate_summary(daily, intervals), indent=2, ensure_ascii=False))
        return

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    daily.to_csv(TABLE_DIR / "table_p2_daily_summary.csv", index=False, float_format="%.10g")
    intervals.to_csv(TABLE_DIR / "table_p2_dispatch.csv", index=False, float_format="%.10g")
    np.savez_compressed(
        TABLE_DIR / "table_p2_scenario_recourse.npz",
        target_dates=daily["date"].to_numpy(dtype="U10"),
        source_residual_dates=source_dates.astype("datetime64[D]"),
        emergency_kw=emergency,
        spill_kw=spill,
        planned_grid_used_kw=planned_grid_used,
        unused_planned_grid_kw=unused_planned_grid,
    )
    summary = _aggregate_summary(daily, intervals)
    (TABLE_DIR / "table_p2_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_table3_outputs(intervals)
    export_result2(daily, intervals)
    print("Q2 expected-cost stochastic MILP complete.")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"result2: {RESULT2_PATH}")


if __name__ == "__main__":
    main()
