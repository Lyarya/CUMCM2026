"""Run the formal Q3 S0-S3 rolling experiment and export official results."""

from __future__ import annotations

import argparse
from concurrent.futures import ProcessPoolExecutor, as_completed
from copy import copy
import hashlib
import json
from pathlib import Path
import shutil
import sys
import tempfile

import numpy as np
import openpyxl
import pandas as pd

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.common.paths import PROJECT_ROOT, RESULTS_DIR
from src.problem2.forecast_interface import (
    FINAL_AUDIT_PATH,
    FORMAL_FORECAST_PATH,
    LEAKAGE_AUDIT_PATH,
    MODEL_DECISION_PATH,
    SCENARIO_MANIFEST_PATH,
    SCENARIO_PATH,
)
from src.problem3.information_schedule import SCHEDULES, SettlementMode
from src.problem3.rolling_dispatch import Q3ScheduleResult, run_schedule


FORMAL_START = "2025-02-01"
FORMAL_END = "2025-12-31"
OUTPUT_DIR = RESULTS_DIR / "problem3"
TABLE_DIR = OUTPUT_DIR / "tables"
TEMPLATE_PATH = PROJECT_ROOT / "data/raw/C题/附件/附件5/result3.xlsx"
RESULT3_PATH = OUTPUT_DIR / "result3.xlsx"
AUDIT_PATH = OUTPUT_DIR / "Q3_ROLLING_AUDIT.md"
SCHEDULE_COMPARISON_PATH = TABLE_DIR / "q3_schedule_comparison.csv"
VOI_PATH = TABLE_DIR / "q3_voi.csv"
DAILY_PATH = TABLE_DIR / "q3_daily_economic_results.csv"
SENSITIVITY_PATH = TABLE_DIR / "q3_settlement_mode_sensitivity.csv"
ARYA_FORECAST_ARTIFACTS = (
    TABLE_DIR / "q3_forecast_audit.json",
    TABLE_DIR / "q3_forecast_fusion_weights.csv",
    TABLE_DIR / "q3_forecast_method_comparison.csv",
    TABLE_DIR / "q3_forecast_predictions.csv",
    TABLE_DIR / "q3_forecast_selection.json",
    TABLE_DIR / "q3_forecast_significance.csv",
    TABLE_DIR / "q3_official_forecast_by_lead.csv",
    TABLE_DIR / "q3_official_forecast_by_release.csv",
    TABLE_DIR / "q3_official_forecast_overall.csv",
)
Q2_PROTECTED_ARTIFACTS = (
    RESULTS_DIR / "problem2/result2.xlsx",
    RESULTS_DIR / "problem2/tables/table_p2_daily_summary.csv",
    RESULTS_DIR / "problem2/tables/table_p2_dispatch.csv",
    RESULTS_DIR / "problem2/tables/table_p2_summary.json",
    FORMAL_FORECAST_PATH,
    SCENARIO_PATH,
    SCENARIO_MANIFEST_PATH,
    MODEL_DECISION_PATH,
    LEAKAGE_AUDIT_PATH,
    FINAL_AUDIT_PATH,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _hash_existing(paths: tuple[Path, ...]) -> dict[str, str]:
    return {str(path): _sha256(path) for path in paths if path.exists()}


def _markdown_table(frame: pd.DataFrame) -> str:
    """Render a compact Markdown table without pandas' optional tabulate dependency."""

    def render(value: object) -> str:
        if isinstance(value, (float, np.floating)):
            return f"{float(value):.6f}"
        return str(value).replace("|", "\\|")

    columns = [str(column) for column in frame.columns]
    lines = [
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in frame.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(render(value) for value in row) + " |")
    return "\n".join(lines)


def _interval_label(slot: int) -> str:
    """Return the interval beginning at (slot-1)*10 minutes."""

    start_minutes = (int(slot) - 1) * 10
    end_minutes = int(slot) * 10
    start_hour, start_minute = divmod(start_minutes, 60)
    end_hour, end_minute = divmod(end_minutes, 60)
    if end_minutes == 24 * 60:
        end_hour = 0
    suffix = "+1" if end_minutes == 24 * 60 else ""
    return f"{start_hour}:{start_minute:02d}-{end_hour}:{end_minute:02d}{suffix}"


def _capture_row_style(sheet, source_row: int, columns: int):
    cells = []
    for column in range(1, columns + 1):
        cell = sheet.cell(source_row, column)
        cells.append(
            (
                copy(cell._style),
                cell.number_format,
                copy(cell.alignment),
                copy(cell.protection),
            )
        )
    return sheet.row_dimensions[source_row].height, cells


def _apply_row_style(sheet, style, row: int) -> None:
    height, cells = style
    sheet.row_dimensions[row].height = height
    for column, (cell_style, number_format, alignment, protection) in enumerate(cells, start=1):
        target = sheet.cell(row, column)
        target._style = copy(cell_style)
        target.number_format = number_format
        target.alignment = copy(alignment)
        target.protection = copy(protection)


def summarize_schedule(result: Q3ScheduleResult) -> dict[str, object]:
    daily = result.daily_main
    intervals = result.intervals
    emergency_count = int((intervals["realized_emergency_kw"] > 1e-9).sum())
    return {
        "schedule": result.schedule,
        "release_hours": ",".join(str(value) for value in SCHEDULES[result.schedule]),
        "risk_setting": "lambda=0 expected-cost",
        "settlement_mode": SettlementMode.SEQUENTIAL_PREVIOUS_COMMITMENT.value,
        "formal_days": int(len(daily)),
        "solver_success": int(daily["solver_success"].sum()),
        "planned_purchase_energy_kwh": float(daily["initial_planned_purchase_energy_kwh"].sum()),
        "planned_purchase_cost_yuan": float(daily["initial_planned_purchase_cost_yuan"].sum()),
        "model_expected_operating_cost_yuan": float(daily["initial_expected_operating_cost_yuan"].sum()),
        "adjustment_energy_kwh": float(daily["adjustment_energy_kwh"].sum()),
        "adjustment_increase_energy_kwh": float(daily["adjustment_increase_energy_kwh"].sum()),
        "adjustment_decrease_energy_kwh": float(daily["adjustment_decrease_energy_kwh"].sum()),
        "adjustment_cost_yuan": float(daily["adjustment_cost_yuan"].sum()),
        "realized_emergency_energy_kwh": float(daily["emergency_energy_kwh"].sum()),
        "realized_emergency_cost_yuan": float(daily["emergency_cost_yuan"].sum()),
        "realized_total_cost_yuan": float(daily["realized_total_cost_yuan"].sum()),
        "emergency_interval_count": emergency_count,
        "emergency_interval_rate": emergency_count / float(len(intervals)),
        "emergency_days": int((daily["emergency_energy_kwh"] > 1e-9).sum()),
        "adjusted_interval_count": int(daily["adjusted_interval_count"].sum()),
        "battery_throughput_kwh": float(daily["battery_throughput_kwh"].sum()),
        "mean_soc_kwh": float(
            intervals[["actual_storage_start_kwh", "actual_storage_end_kwh"]].to_numpy().mean()
        ),
        "minimum_soc_kwh": float(
            intervals[["actual_storage_start_kwh", "actual_storage_end_kwh"]].min().min()
        ),
        "maximum_soc_kwh": float(
            intervals[["actual_storage_start_kwh", "actual_storage_end_kwh"]].max().max()
        ),
        "final_soc_kwh": float(daily["final_energy_kwh"].iloc[-1]),
        "solver_runtime_seconds": float(daily["solver_runtime_seconds"].sum()),
        "maximum_scenario_power_balance_residual_kw": float(
            daily["maximum_scenario_power_balance_residual_kw"].max()
        ),
    }


def summarize_settlement_modes(results: dict[str, Q3ScheduleResult]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for schedule in SCHEDULES:
        frame = results[schedule].daily_sensitivity
        for mode, group in frame.groupby("settlement_mode", sort=False):
            rows.append(
                {
                    "schedule": schedule,
                    "settlement_mode": mode,
                    "formal_days": len(group),
                    "solver_success": int(group["solver_success"].sum()),
                    "planned_purchase_cost_yuan": float(group["initial_planned_purchase_cost_yuan"].sum()),
                    "adjustment_cost_yuan": float(group["adjustment_cost_yuan"].sum()),
                    "realized_emergency_cost_yuan": float(group["emergency_cost_yuan"].sum()),
                    "realized_total_cost_yuan": float(group["realized_total_cost_yuan"].sum()),
                    "adjustment_energy_kwh": float(group["adjustment_energy_kwh"].sum()),
                    "adjustment_increase_energy_kwh": float(group["adjustment_increase_energy_kwh"].sum()),
                    "adjustment_decrease_energy_kwh": float(group["adjustment_decrease_energy_kwh"].sum()),
                    "realized_emergency_energy_kwh": float(group["emergency_energy_kwh"].sum()),
                }
            )
    return pd.DataFrame(rows)


def build_voi(comparison: pd.DataFrame) -> pd.DataFrame:
    costs = comparison.set_index("schedule")["realized_total_cost_yuan"]
    rows = []
    previous = "S0"
    for schedule in SCHEDULES:
        rows.append(
            {
                "schedule": schedule,
                "realized_total_cost_yuan": float(costs[schedule]),
                "voi_vs_s0_yuan": float(costs["S0"] - costs[schedule]),
                "incremental_from_schedule": previous if schedule != "S0" else "S0",
                "incremental_voi_yuan": 0.0 if schedule == "S0" else float(costs[previous] - costs[schedule]),
            }
        )
        previous = schedule
    return pd.DataFrame(rows)


def write_paper_tables(
    comparison: pd.DataFrame,
    voi: pd.DataFrame,
    sensitivity: pd.DataFrame,
) -> tuple[Path, Path, Path]:
    """Generate compact LaTeX tables directly from audited CSV values."""

    comparison_path = TABLE_DIR / "q3_schedule_comparison.tex"
    voi_path = TABLE_DIR / "q3_voi.tex"
    sensitivity_path = TABLE_DIR / "q3_settlement_mode_sensitivity.tex"
    comparison_lines = [
        r"\begin{table}[H]",
        r"\centering\small",
        r"\caption{不同信息安排下的年度实际经济结果}",
        r"\label{tab:q3-schedule-comparison}",
        r"\begin{tabular}{lrrrrr}",
        r"\toprule",
        r"安排 & 计划费用/万元 & 调整费用/万元 & 紧急费用/万元 & 总费用/万元 & 紧急电量/万kWh \\",
        r"\midrule",
    ]
    for row in comparison.itertuples(index=False):
        comparison_lines.append(
            f"{row.schedule} & {row.planned_purchase_cost_yuan / 1e4:.3f} & "
            f"{row.adjustment_cost_yuan / 1e4:.3f} & "
            f"{row.realized_emergency_cost_yuan / 1e4:.3f} & "
            f"{row.realized_total_cost_yuan / 1e4:.3f} & "
            f"{row.realized_emergency_energy_kwh / 1e4:.3f} \\\\"
        )
    comparison_lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}", ""])
    comparison_path.write_text("\n".join(comparison_lines), encoding="utf-8")

    voi_lines = [
        r"\begin{table}[H]",
        r"\centering\small",
        r"\caption{滚动信息的年度价值}",
        r"\label{tab:q3-voi}",
        r"\begin{tabular}{lrr}",
        r"\toprule",
        r"安排 & 相对S0的VOI/元 & 新增发布时间的增量VOI/元 \\",
        r"\midrule",
    ]
    for row in voi.itertuples(index=False):
        voi_lines.append(
            f"{row.schedule} & {row.voi_vs_s0_yuan:.3f} & {row.incremental_voi_yuan:.3f} \\\\"
        )
    voi_lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}", ""])
    voi_path.write_text("\n".join(voi_lines), encoding="utf-8")

    selected = sensitivity.loc[
        sensitivity["schedule"].isin(["S2", "S3"])
    ].copy()
    sensitivity_lines = [
        r"\begin{table}[H]",
        r"\centering\small",
        r"\caption{多次调整结算解释的敏感性结果}",
        r"\label{tab:q3-settlement-sensitivity}",
        r"\begin{tabular}{llrr}",
        r"\toprule",
        r"安排 & 结算解释 & 调整费用/元 & 实际总费用/万元 \\",
        r"\midrule",
    ]
    names = {
        SettlementMode.SEQUENTIAL_PREVIOUS_COMMITMENT.value: "相对上一承诺",
        SettlementMode.ORIGINAL_00_COMMITMENT.value: "相对00:00承诺",
    }
    for row in selected.itertuples(index=False):
        sensitivity_lines.append(
            f"{row.schedule} & {names[row.settlement_mode]} & "
            f"{row.adjustment_cost_yuan:.3f} & {row.realized_total_cost_yuan / 1e4:.3f} \\\\"
        )
    sensitivity_lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}", ""])
    sensitivity_path.write_text("\n".join(sensitivity_lines), encoding="utf-8")
    return comparison_path, voi_path, sensitivity_path


def export_result3(
    daily_s3: pd.DataFrame,
    intervals_s3: pd.DataFrame,
    output_path: Path = RESULT3_PATH,
) -> Path:
    """Fill the official result3 workbook with the formal S3 main settlement."""

    if len(daily_s3) != 334 or len(intervals_s3) != 334 * 144:
        raise ValueError("official result3 export requires 334 days and 144 intervals per day")
    if not intervals_s3.groupby("operating_date")["slot"].nunique().eq(144).all():
        raise ValueError("official result3 export found an incomplete operating day")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_PATH, output_path)
    workbook = openpyxl.load_workbook(output_path)
    expected = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected:
        raise ValueError(f"unexpected result3 sheet structure: {workbook.sheetnames}")

    daily_lookup = daily_s3.set_index("date")
    for sheet_name, power_column, cost_kind in (
        ("计划购电量", "initial_00_grid_kw", "initial"),
        ("调整购电量", "planned_grid_kw", "adjusted"),
    ):
        sheet = workbook[sheet_name]
        if sheet.max_row != 335 or sheet.max_column != 147:
            raise ValueError(f"official {sheet_name} sheet shape changed")
        for slot in range(1, 145):
            sheet.cell(1, slot + 1, _interval_label(slot))
        for day_index, date in enumerate(daily_s3["date"], start=2):
            frame = intervals_s3.loc[intervals_s3["operating_date"].eq(date)].sort_values("slot")
            energy = frame[power_column].to_numpy(float) / 6.0
            for column, value in enumerate(energy, start=2):
                sheet.cell(day_index, column, float(value)).number_format = "0.000000"
            row = daily_lookup.loc[date]
            sheet.cell(day_index, 146, float(energy.sum())).number_format = "0.000000"
            cost = float(row["initial_planned_purchase_cost_yuan"])
            if cost_kind == "adjusted":
                cost += float(row["adjustment_cost_yuan"])
            sheet.cell(day_index, 147, cost).number_format = "0.000000"

    storage = workbook["充放电量"]
    storage_styles = [_capture_row_style(storage, 2 + block, 6) for block in range(6)]
    storage.delete_rows(2, storage.max_row - 1)
    blocks = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    target_row = 2
    for _, day in daily_s3.iterrows():
        frame = intervals_s3.loc[intervals_s3["operating_date"].eq(day["date"])].sort_values("slot")
        for block, label in enumerate(blocks):
            _apply_row_style(storage, storage_styles[block], target_row)
            part = frame.iloc[block * 24 : (block + 1) * 24]
            storage.cell(target_row, 1, pd.Timestamp(day["date"]).to_pydatetime() if block == 0 else None)
            storage.cell(target_row, 2, label)
            storage.cell(target_row, 3, float(part["actual_charge_kwh"].sum())).number_format = "0.000000"
            storage.cell(target_row, 4, float(part["actual_discharge_kwh"].sum())).number_format = "0.000000"
            if block == 0:
                storage.cell(target_row, 5, "00:00")
                storage.cell(target_row, 6, float(day["initial_energy_kwh"])).number_format = "0.000000"
            elif block == 1:
                storage.cell(target_row, 5, "24:00")
                storage.cell(target_row, 6, float(day["final_energy_kwh"])).number_format = "0.000000"
            else:
                storage.cell(target_row, 5, None)
                storage.cell(target_row, 6, None)
            target_row += 1

    emergency = workbook["紧急购电量"]
    emergency_styles = [_capture_row_style(emergency, 2 + index, 3) for index in range(3)]
    emergency.delete_rows(2, emergency.max_row - 1)
    target_row = 2
    for date in daily_s3["date"]:
        frame = intervals_s3.loc[intervals_s3["operating_date"].eq(date)].sort_values("slot")
        for index, row in frame.iterrows():
            style_index = min((target_row - 2) % 144, 2)
            _apply_row_style(emergency, emergency_styles[style_index], target_row)
            emergency.cell(target_row, 1, pd.Timestamp(date).to_pydatetime() if style_index == 0 else None)
            emergency.cell(target_row, 2, _interval_label(int(row["slot"])))
            emergency.cell(target_row, 3, float(row["realized_emergency_kwh"])).number_format = "0.000000"
            target_row += 1

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)
    check = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    if check.sheetnames != expected:
        raise AssertionError("result3 workbook sheet structure changed")
    if check["计划购电量"].max_row != 335 or check["计划购电量"].max_column != 147:
        raise AssertionError("result3 planned-purchase sheet has the wrong shape")
    if check["调整购电量"].max_row != 335 or check["调整购电量"].max_column != 147:
        raise AssertionError("result3 adjusted-purchase sheet has the wrong shape")
    if check["紧急购电量"].max_row != 334 * 144 + 1:
        raise AssertionError("result3 emergency sheet does not contain all intervals")
    return output_path


def write_audit(
    results: dict[str, Q3ScheduleResult],
    comparison: pd.DataFrame,
    voi: pd.DataFrame,
    protected_before: dict[str, str],
    protected_after: dict[str, str],
) -> Path:
    lines = [
        "# Q3 rolling optimization audit",
        "",
        "## Locked contracts",
        "",
        "- Q2 physical baseline: `e86d18d785a6dda28b03d28167c1894b81a100b1`.",
        "- Arya Q3 forecast milestone: `bca6cbfc1c275180023095576954a2e5bfee0bfd` (integrated as cherry-pick `135f7006d359dd7dc2b53442ce319937148480e9`).",
        "- Risk setting: expected cost, lambda = 0.",
        "- Main adjustment settlement: previous commitment.",
        "- Sensitivity settlement: original 00:00 commitment.",
        "- Actual execution calls the unchanged Q2 `settle_realized_day` function.",
        "",
        "## Annual economic results",
        "",
        _markdown_table(comparison),
        "",
        "## Value of information",
        "",
        _markdown_table(voi),
        "",
        "## Physical and causal checks",
        "",
    ]
    for schedule in SCHEDULES:
        lines.extend([f"### {schedule}", ""])
        for key, value in results[schedule].audit.items():
            shown = "PASS" if value is True else "FAIL" if value is False else value
            lines.append(f"- `{key}` = {shown}")
        lines.append("")
    lines.extend(
        [
            "## Artifact protection",
            "",
            f"- Q2 and Arya forecast artifact hashes unchanged: `{protected_before == protected_after}`.",
            "- Future Appendix 2 actual load/PV values are used only by physical execution and ex-post evaluation.",
            "- Scenario count remains 50 and every scenario source date precedes its target date.",
            "- No unabsorbed-discharge or hidden-export accounting channel exists.",
            "",
        ]
    )
    AUDIT_PATH.write_text("\n".join(lines), encoding="utf-8")
    return AUDIT_PATH


def _run_one(schedule: str, dates: pd.DatetimeIndex) -> Q3ScheduleResult:
    checkpoint = Path(tempfile.gettempdir()) / "cumcm2026_q3_checkpoints" / f"{schedule}.pkl"
    return run_schedule(schedule, dates, progress=True, checkpoint_path=checkpoint)


def run_all(dates: pd.DatetimeIndex, workers: int) -> dict[str, Q3ScheduleResult]:
    if workers <= 1:
        return {schedule: _run_one(schedule, dates) for schedule in SCHEDULES}
    results: dict[str, Q3ScheduleResult] = {}
    with ProcessPoolExecutor(max_workers=min(workers, len(SCHEDULES))) as executor:
        futures = {executor.submit(_run_one, schedule, dates): schedule for schedule in SCHEDULES}
        for future in as_completed(futures):
            schedule = futures[future]
            results[schedule] = future.result()
            print(f"Q3 {schedule}: complete", flush=True)
    return {schedule: results[schedule] for schedule in SCHEDULES}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--representative", action="store_true")
    parser.add_argument("--workers", type=int, default=1)
    args = parser.parse_args()
    dates = (
        pd.DatetimeIndex([pd.Timestamp("2025-06-20")])
        if args.representative
        else pd.date_range(FORMAL_START, FORMAL_END, freq="D")
    )
    protected_paths = Q2_PROTECTED_ARTIFACTS + ARYA_FORECAST_ARTIFACTS
    before = _hash_existing(protected_paths)
    results = run_all(dates, args.workers)
    after = _hash_existing(protected_paths)
    if before != after:
        raise AssertionError("Q2 or Arya Q3 protected artifacts changed during the rolling run")
    comparison = pd.DataFrame([summarize_schedule(results[name]) for name in SCHEDULES])
    voi = build_voi(comparison)
    if args.representative:
        print(comparison.to_string(index=False))
        print(voi.to_string(index=False))
        return
    if len(dates) != 334 or not (comparison["solver_success"] == 334).all():
        raise AssertionError("formal Q3 export requires 334/334 Optimal for every schedule")
    for result in results.values():
        if not all(value for value in result.audit.values() if isinstance(value, bool)):
            raise AssertionError(f"{result.schedule} did not pass all Q3 audits")

    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    comparison.to_csv(SCHEDULE_COMPARISON_PATH, index=False, float_format="%.10f")
    voi.to_csv(VOI_PATH, index=False, float_format="%.10f")
    daily = pd.concat(
        [results[name].daily_main.assign(main_settlement=True) for name in SCHEDULES],
        ignore_index=True,
    )
    daily.drop(columns=["wall_runtime_seconds"], errors="ignore").to_csv(
        DAILY_PATH, index=False, float_format="%.10f"
    )
    sensitivity = summarize_settlement_modes(results)
    sensitivity.to_csv(SENSITIVITY_PATH, index=False, float_format="%.10f")
    write_paper_tables(comparison, voi, sensitivity)
    write_audit(results, comparison, voi, before, after)
    export_result3(results["S3"].daily_main, results["S3"].intervals)
    checkpoint_dir = Path(tempfile.gettempdir()) / "cumcm2026_q3_checkpoints"
    for schedule in SCHEDULES:
        checkpoint = checkpoint_dir / f"{schedule}.pkl"
        if checkpoint.exists():
            checkpoint.unlink()
    print(json.dumps({"comparison": comparison.to_dict("records"), "voi": voi.to_dict("records")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
