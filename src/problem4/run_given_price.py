"""Formal full-year GIVEN_PRICE Q4-2/Q4-3 run with resumable checkpoints."""

from __future__ import annotations

import argparse
from concurrent.futures import ProcessPoolExecutor, as_completed
from copy import copy
import json
from pathlib import Path
import shutil
from time import perf_counter

import numpy as np
import openpyxl
import pandas as pd

from src.common.paths import PROJECT_ROOT, RESULTS_DIR
from src.problem2.run import export_result2
from src.problem3.information_schedule import SCHEDULES
from src.problem3.run_rolling import (
    build_voi,
    summarize_schedule,
    summarize_settlement_modes,
)
from src.problem4.dispatch_adapter import GIVEN_PRICE, run_q42_sample, run_q43_sample


FORMAL_DATES = pd.date_range("2025-02-01", "2025-12-31", freq="D")
OUTPUT_DIR = RESULTS_DIR / "problem4"
TABLE_DIR = OUTPUT_DIR / "tables"
CHECKPOINT_DIR = OUTPUT_DIR / "checkpoints" / "given_price"
RESULT42_PATH = OUTPUT_DIR / "result4-2.xlsx"
RESULT43_PATH = OUTPUT_DIR / "result4-3.xlsx"
TEMPLATE42_PATH = PROJECT_ROOT / "data/raw/C题/附件/附件5/result4-2.xlsx"
TEMPLATE43_PATH = PROJECT_ROOT / "data/raw/C题/附件/附件5/result4-3.xlsx"


def _interval_label(slot: int) -> str:
    start = (int(slot) - 1) * 10
    end = int(slot) * 10
    start_hour, start_minute = divmod(start, 60)
    end_hour, end_minute = divmod(end, 60)
    if end == 1440:
        end_hour = 0
    suffix = "+1" if end == 1440 else ""
    return f"{start_hour}:{start_minute:02d}-{end_hour}:{end_minute:02d}{suffix}"


def _correct_result42_labels(path: Path) -> None:
    workbook = openpyxl.load_workbook(path)
    purchase = workbook["计划购电量"]
    for slot in range(1, 145):
        purchase.cell(1, slot + 1, _interval_label(slot))
    emergency = workbook["紧急购电量"]
    for row in range(2, emergency.max_row + 1):
        emergency.cell(row, 2, _interval_label((row - 2) % 144 + 1))
    workbook.save(path)


def _capture_row_style(sheet, source_row: int, columns: int):
    cells = []
    for column in range(1, columns + 1):
        cell = sheet.cell(source_row, column)
        cells.append(
            (
                copy(cell._style),
                cell.number_format,
                copy(cell.alignment),
                copy(cell.protection),
            )
        )
    return sheet.row_dimensions[source_row].height, cells


def _apply_row_style(sheet, style, row: int) -> None:
    height, cells = style
    sheet.row_dimensions[row].height = height
    for column, (cell_style, number_format, alignment, protection) in enumerate(cells, start=1):
        target = sheet.cell(row, column)
        target._style = copy(cell_style)
        target.number_format = number_format
        target.alignment = copy(alignment)
        target.protection = copy(protection)


def _export_result43(daily: pd.DataFrame, intervals: pd.DataFrame) -> None:
    if len(daily) != 334 or len(intervals) != 334 * 144:
        raise ValueError("official result4-3 export requires 334 complete days")
    shutil.copy2(TEMPLATE43_PATH, RESULT43_PATH)
    workbook = openpyxl.load_workbook(RESULT43_PATH)
    expected = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected:
        raise ValueError("result4-3 template sheet structure changed")
    lookup = daily.set_index("date")
    for sheet_name, power_column, adjusted in (
        ("计划购电量", "initial_00_grid_kw", False),
        ("调整购电量", "planned_grid_kw", True),
    ):
        sheet = workbook[sheet_name]
        for slot in range(1, 145):
            sheet.cell(1, slot + 1, _interval_label(slot))
        for target_row, date in enumerate(daily["date"], start=2):
            day = intervals.loc[intervals["operating_date"].eq(date)].sort_values("slot")
            energy = day[power_column].to_numpy(float) / 6.0
            for column, value in enumerate(energy, start=2):
                sheet.cell(target_row, column, float(value)).number_format = "0.000000"
            row = lookup.loc[date]
            sheet.cell(target_row, 146, float(energy.sum())).number_format = "0.000000"
            cost = float(row["initial_planned_purchase_cost_yuan"])
            if adjusted:
                cost += float(row["adjustment_cost_yuan"])
            sheet.cell(target_row, 147, cost).number_format = "0.000000"
    storage = workbook["充放电量"]
    styles = [_capture_row_style(storage, 2 + block, 6) for block in range(6)]
    storage.delete_rows(2, storage.max_row - 1)
    blocks = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    target_row = 2
    for _, day_row in daily.iterrows():
        day = intervals.loc[intervals["operating_date"].eq(day_row["date"])].sort_values("slot")
        for block, label in enumerate(blocks):
            _apply_row_style(storage, styles[block], target_row)
            part = day.iloc[block * 24 : (block + 1) * 24]
            storage.cell(target_row, 1, pd.Timestamp(day_row["date"]).to_pydatetime() if block == 0 else None)
            storage.cell(target_row, 2, label)
            storage.cell(target_row, 3, float(part["actual_charge_kwh"].sum())).number_format = "0.000000"
            storage.cell(target_row, 4, float(part["actual_discharge_kwh"].sum())).number_format = "0.000000"
            if block == 0:
                storage.cell(target_row, 5, "00:00")
                storage.cell(target_row, 6, float(day_row["initial_energy_kwh"])).number_format = "0.000000"
            elif block == 1:
                storage.cell(target_row, 5, "24:00")
                storage.cell(target_row, 6, float(day_row["final_energy_kwh"])).number_format = "0.000000"
            target_row += 1
    emergency = workbook["紧急购电量"]
    styles = [_capture_row_style(emergency, 2 + index, 3) for index in range(3)]
    emergency.delete_rows(2, emergency.max_row - 1)
    target_row = 2
    for date in daily["date"]:
        day = intervals.loc[intervals["operating_date"].eq(date)].sort_values("slot")
        for item_index, (_, row) in enumerate(day.iterrows()):
            _apply_row_style(emergency, styles[min(item_index, 2)], target_row)
            emergency.cell(target_row, 1, pd.Timestamp(date).to_pydatetime() if item_index == 0 else None)
            emergency.cell(target_row, 2, _interval_label(int(row["slot"])))
            emergency.cell(target_row, 3, float(row["realized_emergency_kwh"])).number_format = "0.000000"
            target_row += 1
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(RESULT43_PATH)


def _write_q43_monthly(result, path: Path) -> None:
    frame = result.daily_main.copy()
    frame["month"] = pd.to_datetime(frame["date"]).dt.to_period("M").astype(str)
    monthly = frame.groupby("month", sort=True).agg(
        days_solved=("date", "size"),
        optimal_days=("solver_success", "sum"),
        planned_cost_yuan=("initial_planned_purchase_cost_yuan", "sum"),
        realized_cost_yuan=("realized_total_cost_yuan", "sum"),
        emergency_energy_kwh=("emergency_energy_kwh", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        adjustment_cost_yuan=("adjustment_cost_yuan", "sum"),
        ending_realized_soc_kwh=("final_energy_kwh", "last"),
    ).reset_index()
    monthly.to_csv(path, index=False, float_format="%.10f")


def _workbook_audit(path: Path, *, adjusted: bool) -> dict[str, object]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    expected = (
        ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
        if adjusted
        else ["计划购电量", "充放电量", "紧急购电量"]
    )
    purchase_sheets = ["计划购电量"] + (["调整购电量"] if adjusted else [])
    missing = 0
    invalid = 0
    headers_ok = True
    date_ranges: dict[str, list[str]] = {}
    for name in purchase_sheets:
        sheet = workbook[name]
        headers = [sheet.cell(1, column).value for column in range(2, 146)]
        headers_ok = headers_ok and headers == [_interval_label(slot) for slot in range(1, 145)]
        values = np.array(
            [[sheet.cell(row, column).value for column in range(2, 148)] for row in range(2, 336)],
            dtype=float,
        )
        missing += int(np.isnan(values).sum())
        invalid += int((~np.isfinite(values)).sum())
        dates = [pd.Timestamp(sheet.cell(row, 1).value) for row in range(2, 336)]
        date_ranges[name] = [str(min(dates).date()), str(max(dates).date())]
    emergency = workbook["紧急购电量"]
    emergency_values = np.array(
        [emergency.cell(row, 3).value for row in range(2, emergency.max_row + 1)],
        dtype=float,
    )
    missing += int(np.isnan(emergency_values).sum())
    invalid += int((~np.isfinite(emergency_values)).sum())
    return {
        "path": str(path),
        "sheet_names": workbook.sheetnames,
        "sheet_structure_pass": workbook.sheetnames == expected,
        "planned_sheet_rows": 335,
        "planned_sheet_columns": 147,
        "emergency_rows": emergency.max_row,
        "expected_emergency_rows": 334 * 144 + 1,
        "interval_headers_aligned": headers_ok,
        "timestamp_range": date_ranges,
        "missing_numeric_cells": missing,
        "nonfinite_numeric_cells": invalid,
        "duplicate_interval_labels": 144 - len(set(_interval_label(slot) for slot in range(1, 145))),
        "energy_unit": "kWh",
        "cost_unit": "yuan",
        "numeric_precision": "6 decimal places",
        "pass": bool(
            workbook.sheetnames == expected
            and headers_ok
            and emergency.max_row == 334 * 144 + 1
            and missing == 0
            and invalid == 0
        ),
    }


def _weighted_price(intervals: pd.DataFrame, energy_column: str) -> float | None:
    energy = intervals[energy_column].to_numpy(float)
    total = float(energy.sum())
    if total <= 1e-12:
        return None
    return float(np.sum(intervals["price_yuan_per_kwh"].to_numpy(float) * energy) / total)


def _economic_diagnostics(daily: pd.DataFrame, intervals: pd.DataFrame) -> dict[str, object]:
    costs = daily["realized_total_cost_yuan"].to_numpy(float)
    q1, q3 = np.quantile(costs, [0.25, 0.75])
    upper = float(q3 + 3.0 * (q3 - q1))
    top = daily.nlargest(5, "realized_total_cost_yuan")[["date", "realized_total_cost_yuan"]]
    return {
        "charge_weighted_price_yuan_per_kwh": _weighted_price(intervals, "actual_charge_kwh"),
        "discharge_weighted_price_yuan_per_kwh": _weighted_price(intervals, "actual_discharge_kwh"),
        "high_price_discharge_spread_yuan_per_kwh": (
            (_weighted_price(intervals, "actual_discharge_kwh") or 0.0)
            - (_weighted_price(intervals, "actual_charge_kwh") or 0.0)
        ),
        "daily_cost_outlier_threshold_yuan": upper,
        "daily_cost_outlier_count": int((costs > upper).sum()),
        "top_daily_costs": top.to_dict("records"),
    }


def _summarize_q42(result) -> dict[str, object]:
    daily = result.daily
    intervals = result.intervals
    return {
        "mode": GIVEN_PRICE,
        "formal_days": int(len(daily)),
        "optimal_days": int(daily["solver_status"].eq("Optimal").sum()),
        "planned_purchase_energy_kwh": float(daily["planned_purchase_energy_kwh"].sum()),
        "planned_purchase_cost_yuan": float(daily["planned_purchase_cost_yuan"].sum()),
        "expected_scenario_emergency_cost_yuan": float(daily["expected_scenario_emergency_cost_yuan"].sum()),
        "realized_emergency_energy_kwh": float(daily["realized_emergency_energy_kwh"].sum()),
        "realized_emergency_cost_yuan": float(daily["realized_emergency_cost_yuan"].sum()),
        "realized_total_cost_yuan": float(daily["realized_total_cost_yuan"].sum()),
        "final_soc_kwh": float(daily["actual_final_energy_kwh"].iloc[-1]),
        "solver_runtime_seconds": float(daily["solver_runtime_seconds"].sum()),
        "audit": result.audit,
        "economic_diagnostics": _economic_diagnostics(daily, intervals),
    }


def _run_q43_worker(schedule: str):
    return run_q43_sample(
        schedule,
        FORMAL_DATES,
        mode=GIVEN_PRICE,
        checkpoint_path=CHECKPOINT_DIR / f"q43_{schedule}.pkl",
        progress=True,
    )


def _price_audit(intervals: pd.DataFrame) -> dict[str, object]:
    expected = pd.read_csv(
        TABLE_DIR / "q4_price_forecast_predictions.csv",
        usecols=["operating_date", "slot", "price_yuan_per_kwh"],
    )
    expected = expected.loc[
        expected["operating_date"].astype(str).between("2025-02-01", "2025-12-31")
    ].copy()
    merged = intervals.merge(
        expected,
        on=["operating_date", "slot"],
        how="outer",
        validate="one_to_one",
        suffixes=("_dispatch", "_attachment4"),
        indicator=True,
    )
    residual = abs(merged["price_yuan_per_kwh_dispatch"] - merged["price_yuan_per_kwh_attachment4"])
    emergency_residual = abs(
        intervals["realized_emergency_cost_yuan"]
        - 5.0 * intervals["price_yuan_per_kwh"] * intervals["realized_emergency_kwh"]
    )
    return {
        "one_to_one": bool(merged["_merge"].eq("both").all()),
        "maximum_price_residual_yuan_per_kwh": float(residual.max()),
        "maximum_emergency_cost_residual_yuan": float(emergency_residual.max()),
        "pass": bool(
            merged["_merge"].eq("both").all()
            and residual.max() <= 1e-12
            and emergency_residual.max() <= 2e-3
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workers", type=int, default=4)
    args = parser.parse_args()
    started = perf_counter()
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    CHECKPOINT_DIR.mkdir(parents=True, exist_ok=True)

    q42 = run_q42_sample(
        FORMAL_DATES,
        mode=GIVEN_PRICE,
        checkpoint_path=CHECKPOINT_DIR / "q42.pkl",
        progress=True,
        progress_log_path=TABLE_DIR / "q42_monthly_progress.csv",
    )
    q42.daily.to_csv(TABLE_DIR / "q42_daily.csv", index=False, float_format="%.10f")
    q42.intervals.to_csv(TABLE_DIR / "q42_dispatch.csv", index=False, float_format="%.10f")
    q42_summary = _summarize_q42(q42)
    q42_summary["price_alignment"] = _price_audit(q42.intervals)
    (TABLE_DIR / "q42_summary.json").write_text(
        json.dumps(q42_summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    q42_daily_export = q42.daily.copy()
    q42_daily_export["actual_initial_energy_kwh"] = q42_daily_export["initial_energy_kwh"]
    export_result2(
        q42_daily_export,
        q42.intervals,
        output_path=RESULT42_PATH,
        template_path=TEMPLATE42_PATH,
    )
    _correct_result42_labels(RESULT42_PATH)
    print("Q4-2 GIVEN_PRICE complete and audited", flush=True)

    schedules: dict[str, object] = {}
    if args.workers <= 1:
        schedules = {name: _run_q43_worker(name) for name in SCHEDULES}
    else:
        with ProcessPoolExecutor(max_workers=min(args.workers, len(SCHEDULES))) as executor:
            futures = {executor.submit(_run_q43_worker, name): name for name in SCHEDULES}
            for future in as_completed(futures):
                name = futures[future]
                schedules[name] = future.result()
                print(f"Q4-3 {name} GIVEN_PRICE complete", flush=True)
        schedules = {name: schedules[name] for name in SCHEDULES}

    comparison = pd.DataFrame([summarize_schedule(schedules[name]) for name in SCHEDULES])
    voi = build_voi(comparison)
    sensitivity = summarize_settlement_modes(schedules)
    comparison.to_csv(TABLE_DIR / "q43_schedule_comparison.csv", index=False, float_format="%.10f")
    voi.to_csv(TABLE_DIR / "q43_voi.csv", index=False, float_format="%.10f")
    sensitivity.to_csv(TABLE_DIR / "q43_settlement_sensitivity.csv", index=False, float_format="%.10f")
    pd.concat(
        [schedules[name].daily_main for name in SCHEDULES], ignore_index=True
    ).to_csv(TABLE_DIR / "q43_daily.csv", index=False, float_format="%.10f")
    for name in SCHEDULES:
        _write_q43_monthly(
            schedules[name], TABLE_DIR / f"q43_{name}_monthly_progress.csv"
        )
    _export_result43(schedules["S3"].daily_main, schedules["S3"].intervals)

    q43_audits = {
        name: {
            "physical": schedules[name].audit,
            "price_alignment": _price_audit(schedules[name].intervals),
            "economic_diagnostics": _economic_diagnostics(
                schedules[name].daily_main, schedules[name].intervals
            ),
        }
        for name in SCHEDULES
    }
    workbook_audits = {
        "result4-2": _workbook_audit(RESULT42_PATH, adjusted=False),
        "result4-3": _workbook_audit(RESULT43_PATH, adjusted=True),
    }
    fixed_q2 = json.loads(
        (RESULTS_DIR / "problem2/tables/table_p2_summary.json").read_text(encoding="utf-8")
    )
    fixed_q3 = pd.read_csv(RESULTS_DIR / "problem3/tables/q3_schedule_comparison.csv")
    economic = {
        "q2_fixed_price_realized_total_cost_yuan": fixed_q2["realized_total_cost_yuan"],
        "q42_dynamic_price_realized_total_cost_yuan": q42_summary["realized_total_cost_yuan"],
        "q3_fixed_price_by_schedule": fixed_q3.set_index("schedule")["realized_total_cost_yuan"].to_dict(),
        "q43_dynamic_price_by_schedule": comparison.set_index("schedule")["realized_total_cost_yuan"].to_dict(),
        "best_q43_schedule": str(comparison.loc[comparison["realized_total_cost_yuan"].idxmin(), "schedule"]),
        "incremental_voi": voi.set_index("schedule")["incremental_voi_yuan"].to_dict(),
    }
    audit = {
        "mode": GIVEN_PRICE,
        "formal_dates": [str(FORMAL_DATES[0].date()), str(FORMAL_DATES[-1].date())],
        "q42": q42_summary,
        "q43": q43_audits,
        "workbooks": workbook_audits,
        "economic_sanity": economic,
        "runtime_seconds": perf_counter() - started,
        "all_pass": bool(
            q42_summary["optimal_days"] == 334
            and all(row == 334 for row in comparison["solver_success"])
            and all(value for value in q42.audit.values() if isinstance(value, bool))
            and all(
                all(value for value in schedules[name].audit.values() if isinstance(value, bool))
                and q43_audits[name]["price_alignment"]["pass"]
                for name in SCHEDULES
            )
            and all(value["pass"] for value in workbook_audits.values())
        ),
    }
    (OUTPUT_DIR / "Q4_GIVEN_PRICE_FULL_AUDIT.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not audit["all_pass"]:
        raise AssertionError("Q4 GIVEN_PRICE annual audit failed; inspect audit JSON")
    print(json.dumps(audit, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
