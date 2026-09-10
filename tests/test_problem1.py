"""End-to-end checks for the deterministic Problem 1 dispatch."""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from src.common.paths import PROCESSED_DATA_DIR
from src.problem1.evaluate import compute_baseline, validate_dispatch
from src.problem1.model import DispatchParameters, solve_deterministic_dispatch
from src.problem1.run import TEMPLATE_PATH, export_result1


@pytest.fixture(scope="module")
def solved_dispatch():
    data = pd.read_csv(PROCESSED_DATA_DIR / "C题" / "problem1_day.csv")
    parameters = DispatchParameters()
    result = solve_deterministic_dispatch(data, parameters)
    validation = validate_dispatch(result, parameters)
    return data, result, validation


def test_problem1_solution_is_feasible_and_improves_cost(solved_dispatch) -> None:
    data, result, validation = solved_dispatch
    baseline = compute_baseline(data)

    assert result.status == "Optimal"
    assert result.objective_yuan < baseline["baseline_cost_yuan"]
    assert validation.period_count == 144
    assert validation.initial_energy_kwh == pytest.approx(6_000.0, abs=1e-5)
    assert validation.terminal_energy_kwh == pytest.approx(6_000.0, abs=1e-5)
    assert validation.minimum_energy_kwh >= 1_200.0 - 1e-5
    assert validation.maximum_energy_kwh <= 10_800.0 + 1e-5
    assert validation.simultaneous_charge_discharge_count == 0
    assert validation.maximum_power_balance_residual_kw <= 1e-3
    assert validation.maximum_energy_transition_residual_kwh <= 1e-3


def test_result1_preserves_official_template(solved_dispatch, tmp_path) -> None:
    _, result, _ = solved_dispatch
    output_path = tmp_path / "result1.xlsx"
    export_result1(result.dispatch, output_path=output_path)

    template = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    workbook = openpyxl.load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == template.sheetnames == ["计划购电量", "充放电量"]
    assert workbook["计划购电量"].max_row == 145
    assert workbook["计划购电量"].max_column == 2
    assert all(
        workbook["计划购电量"].cell(row=row, column=2).value is not None
        for row in range(2, 146)
    )
    assert workbook["充放电量"]["E2"].value == pytest.approx(6_000.0)
    assert workbook["充放电量"]["E3"].value == pytest.approx(6_000.0)
