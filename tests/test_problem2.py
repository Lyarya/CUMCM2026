"""Physical and accounting tests for the first Q2 expected-cost MILP."""

from __future__ import annotations

import hashlib

import numpy as np
import openpyxl
import pandas as pd
import pytest

from src.problem2.evaluate import validate_q2_day
from src.problem2.forecast_interface import SCENARIO_PATH, get_q2_day_inputs
from src.problem2.model import Q2DispatchParameters, solve_expected_cost_dispatch
from src.problem2.run import (
    RESULT2_PATH,
    TABLE_DIR,
    _interval_label,
    run_chronological,
    settle_realized_day,
)


def _synthetic_day(
    *, load_kw: float = 0.0, pv_kw: float = 0.0, grid_kw: float = 0.0,
    charge_kw: float = 0.0, discharge_kw: float = 0.0,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    timestamps = pd.date_range("2025-01-01 00:10", periods=144, freq="10min")
    plan = pd.DataFrame(
        {
            "slot": np.arange(1, 145),
            "timestamp": timestamps,
            "planned_grid_kw": grid_kw,
            "planned_charge_limit_kw": charge_kw,
            "planned_discharge_limit_kw": discharge_kw,
            "price_yuan_per_kwh": 1.0,
        }
    )
    actual = pd.DataFrame(
        {
            "datetime": timestamps,
            "actual_load": load_kw,
            "actual_generation": pv_kw,
        }
    )
    return plan, actual


def _sha256(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def representative_solution():
    inputs = get_q2_day_inputs("2025-07-15", initial_energy=6_000.0)
    result = solve_expected_cost_dispatch(inputs)
    report = validate_q2_day(inputs, result)
    return inputs, result, report


def test_expected_cost_milp_is_optimal_and_finite(representative_solution) -> None:
    inputs, result, report = representative_solution
    assert result.status == "Optimal"
    assert result.solver_name == "PuLP HiGHS"
    assert report.period_count == 144
    assert report.scenario_count == 50
    assert result.emergency_kw.shape == (50, 144)
    assert result.spill_kw.shape == (50, 144)
    assert result.planned_grid_used_kw.shape == (50, 144)
    assert result.unused_planned_grid_kw.shape == (50, 144)
    assert np.isfinite(result.dispatch.select_dtypes(include=[np.number])).all().all()
    assert np.isfinite(result.emergency_kw).all()
    assert np.isfinite(result.spill_kw).all()
    assert np.isfinite(result.planned_grid_used_kw).all()
    assert np.isfinite(result.unused_planned_grid_kw).all()
    assert inputs.dt_hours == pytest.approx(1 / 6)


def test_q2_physics_pairing_and_five_times_price(representative_solution) -> None:
    inputs, result, report = representative_solution
    parameters = Q2DispatchParameters()
    expected_emergency_cost = (
        parameters.emergency_price_multiplier
        * np.mean(
            np.sum(inputs.price[None, :] * result.emergency_kw * inputs.dt_hours, axis=1)
        )
    )
    assert report.scenario_pairs_unchanged is True
    assert report.simultaneous_charge_discharge_count == 0
    assert report.negative_grid_purchase_count == 0
    assert report.negative_emergency_purchase_count == 0
    assert report.negative_planned_grid_use_count == 0
    assert report.maximum_planned_grid_use_excess_kw <= 2e-3
    assert report.maximum_unused_grid_identity_residual_kw <= 2e-3
    assert report.maximum_spill_excess_kw <= 2e-3
    assert report.maximum_power_balance_residual_kw <= 2e-3
    assert report.maximum_energy_transition_residual_kwh <= 2e-3
    assert report.minimum_energy_kwh >= parameters.minimum_energy_kwh - 2e-3
    assert report.maximum_energy_kwh <= parameters.maximum_energy_kwh + 2e-3
    assert result.expected_emergency_cost_yuan == pytest.approx(expected_emergency_cost, abs=2e-3)


def test_solver_does_not_modify_frozen_scenarios(representative_solution) -> None:
    before = _sha256(SCENARIO_PATH)
    inputs = get_q2_day_inputs("2025-10-15", initial_energy=7_321.5)
    load_before = inputs.load_scenarios.copy()
    pv_before = inputs.pv_scenarios.copy()
    solve_expected_cost_dispatch(inputs)
    assert np.array_equal(inputs.load_scenarios, load_before)
    assert np.array_equal(inputs.pv_scenarios, pv_before)
    assert _sha256(SCENARIO_PATH) == before


def test_cross_day_soc_is_carried_without_reset_to_6000() -> None:
    daily, intervals, emergency, spill, planned_used, unused_plan, sources = run_chronological(
        pd.date_range("2025-02-01", "2025-02-02", freq="D"),
        initial_energy_kwh=7_321.5,
    )
    assert daily.loc[0, "initial_energy_kwh"] == pytest.approx(7_321.5)
    assert daily.loc[0, "initial_energy_kwh"] != pytest.approx(6_000.0)
    assert daily.loc[1, "initial_energy_kwh"] == pytest.approx(
        daily.loc[0, "final_energy_kwh"], abs=2e-3
    )
    assert abs(daily.loc[0, "final_energy_kwh"] - daily.loc[0, "initial_energy_kwh"]) > 2e-3
    assert daily["final_energy_kwh"].min() > 1_200.0 + 2e-3
    assert intervals.shape[0] == 288
    assert emergency.shape == (2, 50, 144)
    assert spill.shape == (2, 50, 144)
    assert planned_used.shape == (2, 50, 144)
    assert unused_plan.shape == (2, 50, 144)
    assert sources.shape == (2, 50)


def test_full_period_outputs_preserve_result2_template() -> None:
    daily = pd.read_csv(TABLE_DIR / "table_p2_daily_summary.csv")
    intervals = pd.read_csv(TABLE_DIR / "table_p2_dispatch.csv")
    assert len(daily) == 334
    assert len(intervals) == 334 * 144
    assert (daily["status"] == "Optimal").all()
    assert np.max(
        np.abs(
            daily["initial_energy_kwh"].iloc[1:].to_numpy()
            - daily["final_energy_kwh"].iloc[:-1].to_numpy()
        )
    ) <= 2e-3

    with np.load(TABLE_DIR / "table_p2_scenario_recourse.npz", allow_pickle=False) as archive:
        assert archive["emergency_kw"].shape == (334, 50, 144)
        assert archive["spill_kw"].shape == (334, 50, 144)
        assert archive["planned_grid_used_kw"].shape == (334, 50, 144)
        assert archive["unused_planned_grid_kw"].shape == (334, 50, 144)
        assert archive["source_residual_dates"].shape == (334, 50)
        assert np.isfinite(archive["emergency_kw"]).all()
        assert np.isfinite(archive["spill_kw"]).all()
        assert (archive["emergency_kw"] >= -2e-3).all()
        assert (archive["spill_kw"] >= -2e-3).all()
        assert (archive["planned_grid_used_kw"] >= -2e-3).all()
        assert (archive["unused_planned_grid_kw"] >= -2e-3).all()

    workbook = openpyxl.load_workbook(RESULT2_PATH, read_only=True, data_only=True)
    assert workbook.sheetnames == ["计划购电量", "充放电量", "紧急购电量"]
    assert (workbook["计划购电量"].max_row, workbook["计划购电量"].max_column) == (335, 147)
    assert (workbook["充放电量"].max_row, workbook["充放电量"].max_column) == (2005, 6)
    assert (workbook["紧急购电量"].max_row, workbook["紧急购电量"].max_column) == (334 * 144 + 1, 3)
    assert workbook["计划购电量"]["A2"].value.date().isoformat() == "2025-02-01"
    assert workbook["计划购电量"]["A335"].value.date().isoformat() == "2025-12-31"
    assert workbook["紧急购电量"]["B2"].value == "0:10-0:20"
    assert workbook["紧急购电量"]["B145"].value == "0:00-0:10+1"

    table3 = pd.read_csv(TABLE_DIR / "table_p2_table3_emergency.csv")
    assert table3.shape == (144, 8)
    assert list(table3.columns) == [
        "2025.3.20_时间段",
        "2025.3.20_购电量_kWh",
        "2025.6.21_时间段",
        "2025.6.21_购电量_kWh",
        "2025.9.23_时间段",
        "2025.9.23_购电量_kWh",
        "2025.12.21_时间段",
        "2025.12.21_购电量_kWh",
    ]
    assert not table3.isna().any().any()
    assert (table3.filter(like="购电量_kWh") >= 0.0).all().all()
    assert (table3.filter(like="购电量_kWh") == 0.0).any().all()
    assert table3["2025.3.20_时间段"].iloc[0] == "0:10-0:20"
    assert table3["2025.3.20_时间段"].iloc[-1] == "0:00-0:10+1"


def test_official_interval_labels_follow_template() -> None:
    assert _interval_label(1) == "0:10-0:20"
    assert _interval_label(143) == "23:50-0:00+1"
    assert _interval_label(144) == "0:00-0:10+1"


def test_realized_settlement_has_no_unexplained_discharge_sink() -> None:
    plan, actual = _synthetic_day(load_kw=20.0, discharge_kw=100.0)
    settled, final_energy = settle_realized_day(plan, actual, 6_000.0)
    assert "realized_unabsorbed_discharge_kw" not in settled.columns
    assert settled["actual_discharge_kw"].max() <= 20.0 + 1e-9
    balance = (
        settled["realized_planned_grid_used_kw"]
        + settled["realized_emergency_kw"]
        + settled["actual_pv_kw"]
        + settled["actual_discharge_kw"]
        - settled["actual_load_kw"]
        - settled["actual_charge_kw"]
        - settled["realized_pv_spill_kw"]
    )
    assert balance.abs().max() < 1e-9
    assert final_energy >= 1_200.0


def test_realized_settlement_covers_unused_grid_pv_shortage_and_soc_bounds() -> None:
    plan, actual = _synthetic_day(load_kw=50.0, grid_kw=100.0)
    settled, _ = settle_realized_day(plan, actual, 6_000.0)
    assert settled["realized_unused_planned_grid_kwh"].sum() == pytest.approx(1_200.0)
    assert settled["realized_emergency_kwh"].sum() == pytest.approx(0.0)

    plan, actual = _synthetic_day(load_kw=10.0, pv_kw=100.0)
    settled, _ = settle_realized_day(plan, actual, 6_000.0)
    assert settled["realized_pv_spill_kwh"].sum() == pytest.approx(2_160.0)
    assert (settled["realized_pv_spill_kw"] <= settled["actual_pv_kw"] + 1e-9).all()

    plan, actual = _synthetic_day(load_kw=200.0, grid_kw=100.0)
    settled, _ = settle_realized_day(plan, actual, 6_000.0)
    assert settled["realized_emergency_kwh"].sum() == pytest.approx(2_400.0)
    assert settled["realized_emergency_cost_yuan"].sum() == pytest.approx(12_000.0)

    plan, actual = _synthetic_day(charge_kw=5_000.0)
    settled, final_energy = settle_realized_day(plan, actual, 10_799.0)
    assert final_energy <= 10_800.0 + 1e-9
    plan, actual = _synthetic_day(load_kw=5_000.0, discharge_kw=5_000.0)
    settled, final_energy = settle_realized_day(plan, actual, 1_201.0)
    assert final_energy >= 1_200.0 - 1e-9
