"""Physical, causal, and accounting tests for Q3 rolling optimization."""

from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from src.problem2.forecast_interface import (
    FORMAL_FORECAST_PATH,
    SCENARIO_PATH,
    get_q2_day_inputs,
)
from src.problem3.information_schedule import SettlementMode
from src.problem3.rolling_dispatch import (
    build_release_inputs,
    run_schedule,
)
from src.problem3.rolling_model import solve_remaining_dispatch
from src.problem3.visualize import summarize_economic_results


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FORECAST_ARTIFACTS = (
    PROJECT_ROOT / "results/problem3/tables/q3_forecast_predictions.csv",
    PROJECT_ROOT / "results/problem3/tables/q3_forecast_selection.json",
    PROJECT_ROOT / "results/problem3/tables/q3_forecast_audit.json",
)


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def representative_s3():
    return run_schedule(
        "S3", pd.DatetimeIndex([pd.Timestamp("2025-06-20")]), progress=False
    )


def test_release_inputs_keep_locked_scenarios_and_are_causal() -> None:
    before = _hash(SCENARIO_PATH)
    inputs, update = build_release_inputs("2025-06-20", 12, 5432.1)
    base = get_q2_day_inputs("2025-06-20", initial_energy=5432.1)
    assert inputs.horizon == 72
    assert inputs.scenario_count == 50
    assert inputs.initial_energy == pytest.approx(5432.1)
    np.testing.assert_array_equal(inputs.load_scenarios, base.load_scenarios[:, 72:])
    np.testing.assert_array_equal(inputs.scenario_source_dates, base.scenario_source_dates)
    residual = base.pv_scenarios[:, 72:] - base.pv_forecast[None, 72:]
    np.testing.assert_allclose(
        inputs.pv_scenarios,
        np.maximum(inputs.pv_forecast[None, :] + residual, 0.0),
        atol=1e-10,
    )
    assert (inputs.scenario_source_dates < np.datetime64("2025-06-20")).all()
    assert update.executed_mask.sum() == 72
    assert update.future_mask.sum() == 72
    assert _hash(SCENARIO_PATH) == before


def test_remaining_horizon_model_uses_expected_cost_and_balances_scenarios() -> None:
    inputs, _ = build_release_inputs("2025-06-20", 18, 6000.0)
    result = solve_remaining_dispatch(inputs)
    assert result.status == "Optimal"
    assert len(result.dispatch) == 36
    assert result.emergency_kw.shape == (50, 36)
    assert result.maximum_scenario_power_balance_residual_kw <= 2e-3
    assert result.operating_cost_yuan == pytest.approx(
        result.planned_purchase_cost_yuan + result.expected_emergency_cost_yuan
    )
    assert result.optimization_objective_yuan == pytest.approx(
        result.operating_cost_yuan - result.terminal_value_credit_yuan
    )


def test_revision_model_uses_incremental_adjustment_prices() -> None:
    inputs, _ = build_release_inputs("2025-06-20", 18, 6000.0)
    prior = np.full(inputs.horizon, 3000.0)
    result = solve_remaining_dispatch(inputs, prior_commitment_kw=prior)
    grid = result.dispatch["planned_grid_kw"].to_numpy(float)
    delta_kwh = (grid - prior) * inputs.dt_hours
    expected_adjustment = np.sum(
        inputs.price
        * (1.5 * np.maximum(delta_kwh, 0.0) - 0.5 * np.maximum(-delta_kwh, 0.0))
    )
    assert result.adjustment_cost_yuan == pytest.approx(expected_adjustment, abs=1e-5)
    assert result.optimization_objective_yuan == pytest.approx(
        result.adjustment_cost_yuan
        + result.expected_emergency_cost_yuan
        - result.terminal_value_credit_yuan,
        abs=1e-5,
    )


def test_solver_tolerance_negative_commitment_is_normalized() -> None:
    inputs, _ = build_release_inputs("2025-06-20", 18, 6000.0)
    tolerated = np.full(inputs.horizon, 3000.0)
    tolerated[0] = -1e-8
    result = solve_remaining_dispatch(inputs, prior_commitment_kw=tolerated)
    assert (result.dispatch["planned_grid_kw"] >= 0.0).all()
    rejected = tolerated.copy()
    rejected[0] = -1e-3
    with pytest.raises(ValueError, match="nonnegative"):
        solve_remaining_dispatch(inputs, prior_commitment_kw=rejected)


def test_s3_representative_day_passes_physical_and_causal_audit(representative_s3) -> None:
    result = representative_s3
    assert len(result.intervals) == 144
    assert result.intervals["slot"].tolist() == list(range(1, 145))
    assert set(result.intervals["release_hour"]) == {0, 6, 12, 18}
    assert result.intervals.groupby("release_hour").size().to_dict() == {
        0: 36, 6: 36, 12: 36, 18: 36
    }
    assert all(value for value in result.audit.values() if isinstance(value, bool))
    assert result.audit["unabsorbed_discharge_energy_kwh"] == 0.0
    assert result.daily_main.iloc[0]["release_count"] == 4


def test_two_settlement_modes_reconcile_without_double_counting(representative_s3) -> None:
    daily = representative_s3.daily_sensitivity
    assert set(daily["settlement_mode"]) == {
        SettlementMode.SEQUENTIAL_PREVIOUS_COMMITMENT.value,
        SettlementMode.ORIGINAL_00_COMMITMENT.value,
    }
    expected = (
        daily["initial_planned_purchase_cost_yuan"]
        + daily["adjustment_cost_yuan"]
        + daily["emergency_cost_yuan"]
    )
    np.testing.assert_allclose(daily["realized_total_cost_yuan"], expected, atol=1e-6)
    assert daily["emergency_cost_yuan"].nunique() == 1
    assert daily["initial_planned_purchase_cost_yuan"].nunique() == 1


def test_q2_and_arya_forecast_artifacts_are_not_modified(representative_s3) -> None:
    paths = (SCENARIO_PATH, FORMAL_FORECAST_PATH, *FORECAST_ARTIFACTS)
    assert all(path.exists() for path in paths)
    before = {path: _hash(path) for path in paths}
    _ = representative_s3.audit
    after = {path: _hash(path) for path in paths}
    assert before == after


def test_formal_result3_matches_audited_s3_costs_and_time_grid() -> None:
    result_path = PROJECT_ROOT / "results/problem3/result3.xlsx"
    comparison_path = PROJECT_ROOT / "results/problem3/tables/q3_schedule_comparison.csv"
    daily_path = PROJECT_ROOT / "results/problem3/tables/q3_daily_economic_results.csv"
    assert result_path.exists()
    comparison = pd.read_csv(comparison_path).set_index("schedule").loc["S3"]
    daily = pd.read_csv(daily_path)
    assert daily.groupby("schedule").size().to_dict() == {
        "S0": 334, "S1": 334, "S2": 334, "S3": 334
    }
    workbook = openpyxl.load_workbook(result_path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    plan = workbook["计划购电量"]
    adjusted = workbook["调整购电量"]
    emergency = workbook["紧急购电量"]
    assert (plan.max_row, plan.max_column) == (335, 147)
    assert (adjusted.max_row, adjusted.max_column) == (335, 147)
    assert (emergency.max_row, emergency.max_column) == (334 * 144 + 1, 3)
    assert plan.cell(1, 2).value == "0:00-0:10"
    assert plan.cell(1, 145).value == "23:50-0:00+1"
    assert adjusted.cell(1, 2).value == "0:00-0:10"
    assert adjusted.cell(1, 145).value == "23:50-0:00+1"
    plan_rows = list(plan.iter_rows(min_row=2, values_only=True))
    adjusted_rows = list(adjusted.iter_rows(min_row=2, values_only=True))
    plan_energy = sum(float(row[145]) for row in plan_rows)
    plan_cost = sum(float(row[146]) for row in plan_rows)
    adjusted_cost = sum(float(row[146]) for row in adjusted_rows)
    emergency_energy = sum(
        float(row[2]) for row in emergency.iter_rows(min_row=2, values_only=True)
    )
    assert plan_energy == pytest.approx(comparison["planned_purchase_energy_kwh"], abs=1e-4)
    assert plan_cost == pytest.approx(comparison["planned_purchase_cost_yuan"], abs=1e-4)
    assert adjusted_cost == pytest.approx(
        comparison["planned_purchase_cost_yuan"] + comparison["adjustment_cost_yuan"],
        abs=1e-4,
    )
    assert emergency_energy == pytest.approx(
        comparison["realized_emergency_energy_kwh"], abs=1e-4
    )


def test_formal_result3_emergency_cost_uses_five_times_fixed_tariff() -> None:
    workbook = openpyxl.load_workbook(
        PROJECT_ROOT / "results/problem3/result3.xlsx", read_only=True, data_only=True
    )
    emergency = workbook["紧急购电量"]
    price = get_q2_day_inputs("2025-02-01", initial_energy=6000.0).price
    energy = np.fromiter(
        (float(row[2]) for row in emergency.iter_rows(min_row=2, values_only=True)),
        dtype=float,
        count=334 * 144,
    ).reshape(334, 144)
    recomputed = float(np.sum(energy * price[None, :] * 5.0))
    comparison = pd.read_csv(
        PROJECT_ROOT / "results/problem3/tables/q3_schedule_comparison.csv"
    ).set_index("schedule").loc["S3"]
    assert recomputed == pytest.approx(
        comparison["realized_emergency_cost_yuan"], abs=1e-4
    )
    assert (
        comparison["planned_purchase_cost_yuan"]
        + comparison["adjustment_cost_yuan"]
        + recomputed
    ) == pytest.approx(comparison["realized_total_cost_yuan"], abs=1e-4)


def test_final_economic_reporting_is_derived_from_canonical_tables() -> None:
    table_dir = PROJECT_ROOT / "results/problem3/tables"
    schedule = pd.read_csv(table_dir / "q3_schedule_comparison.csv")
    voi = pd.read_csv(table_dir / "q3_voi.csv")
    settlement = pd.read_csv(table_dir / "q3_settlement_mode_sensitivity.csv")
    summary = summarize_economic_results(schedule, voi, settlement)

    assert summary["best_schedule"] == "S2"
    assert summary["savings_vs_s0_yuan"] == pytest.approx(222469.864818, abs=1e-6)
    assert summary["savings_vs_s0_pct"] == pytest.approx(1.51052626946)
    assert summary["emergency_reduction_vs_s0_pct"] == pytest.approx(17.4946354238)
    assert summary["incremental_voi_yuan"]["S3"] == pytest.approx(-2217.412836, abs=1e-6)
    assert summary["settlement_ranking_stable"] is True
    assert set(map(tuple, summary["settlement_rankings"].values())) == {
        ("S2", "S3", "S1", "S0")
    }
